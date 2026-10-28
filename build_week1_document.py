"""Build WEEK1_DEVELOPMENT_DOCUMENT.docx from the Markdown source.

The Markdown in docs/ is the editable source; this script applies the black and
dark-blue theme, cover page, contents, styled tables and page numbering.

Usage (from the project root):
    pip install -r requirements-dev.txt
    python scripts/build_week1_document.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "WEEK1_DEVELOPMENT_DOCUMENT.md"
OUTPUT = ROOT / "WEEK1_DEVELOPMENT_DOCUMENT.docx"

INK = RGBColor(0x0B, 0x0B, 0x0C)
NAVY = RGBColor(0x0F, 0x2C, 0x59)
MID_BLUE = RGBColor(0x1F, 0x4E, 0x8C)
SOFT_BLUE = RGBColor(0x8F, 0xB4, 0xE3)
PALE_BLUE = RGBColor(0xD6, 0xE4, 0xF5)
GREY = RGBColor(0x5A, 0x62, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NAVY_HEX = "0F2C59"
BLACK_HEX = "0B0B0C"
CARD_HEX = "F4F6F9"
NOTE_HEX = "EEF2F8"
CODE_HEX = "F5F6F8"
BORDER_HEX = "D5DCE6"

TEXT_WIDTH_CM = 16.6
USABLE_WIDTH = Cm(TEXT_WIDTH_CM)

META = {
    "prepared_by": "Sumyai Butt",
    "project": "AI Customer Calling Agent",
    "document_type": "Week 1 Development Progress Document",
    "week": "Week 1 - Discovery and Foundation (Days 1 to 5)",
    "environment": "Python 3.12.10 on Windows (win32)",
    "related": "AI_Customer_Calling_Agent_Development_Plan.docx (unchanged)",
    "repository": "New Week 1 files only; nothing overwritten or deleted",
}

# Column widths in percent, keyed by the table's header row. Falls back to a
# sensible default for any table not listed here.
WIDTH_MAP = {
    ("Item", "Detail"): [28, 72],
    ("Week 1 requirement", "Status", "Evidence"): [36, 26, 38],
    ("Intent", "Example customer phrases", "Agent behaviour"): [18, 40, 42],
    ("In scope (implemented)", "Out of scope (planned later)"): [50, 50],
    ("ID", "Requirement", "Priority", "Status"): [12, 50, 14, 24],
    ("Layer", "Responsibility", "Week 1 status"): [22, 54, 24],
    ("Stage", "Component", "What it does"): [7, 28, 65],
    ("Component", "Week 1 status", "Planned week"): [38, 30, 32],
    ("Step", "Master plan call flow", "Week 1 implementation"): [8, 34, 58],
    ("Order", "Rule group", "Reason for this position"): [8, 20, 72],
    ("Situation", "Week 1 behaviour", "Escalation flag"): [38, 44, 18],
    ("Speaker", "Message"): [18, 82],
    ("File", "Responsibility", "Key functions"): [28, 42, 30],
    ("Field", "Type", "Meaning"): [20, 14, 66],
    ("Test area", "Tests", "What is verified"): [38, 10, 52],
    ("ID", "Scenario", "Input", "Expected result"): [9, 19, 36, 36],
    ("Item", "Result"): [30, 70],
    ("Defect", "Cause", "Fix", "Verification"): [30, 30, 28, 12],
    ("Access or decision", "Purpose", "Status"): [38, 48, 14],
    ("Day", "Planned task", "Actual output", "Status"): [8, 25, 47, 20],
    ("Risk", "Likelihood", "Impact", "Mitigation"): [34, 13, 10, 43],
    ("Deliverable", "Status", "Evidence"): [32, 24, 44],
    ("Limitation", "Impact", "Planned resolution"): [34, 34, 32],
    ("File", "Type", "Purpose"): [32, 10, 58],
    ("Term", "Meaning"): [20, 80],
    ("Check", "Result"): [35, 65],
}

DEFAULT_RATIOS = {
    2: [36, 64],
    3: [24, 38, 38],
    4: [12, 36, 20, 32],
}


# ---------------------------------------------------------------------------
# Low-level formatting helpers
# ---------------------------------------------------------------------------

def style_run(
    run,
    name: str = "Calibri",
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        rfonts.set(qn("w:eastAsia"), name)
        rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def shade_cell(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)


def shade_paragraph(paragraph, hex_fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)


def add_paragraph_border(
    paragraph, edge: str = "bottom", color: str = NAVY_HEX, size: int = 8, space: int = 4
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    p_bdr.append(element)


def set_table_borders(
    table,
    color: str = BORDER_HEX,
    outer: str = NAVY_HEX,
    vertical: bool = False,
) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = [
        ("top", "single", 10, outer),
        ("bottom", "single", 10, outer),
        ("left", "none", 0, "auto"),
        ("right", "none", 0, "auto"),
        ("insideH", "single", 4, color),
        ("insideV", "single" if vertical else "none", 4 if vertical else 0, color),
    ]
    for edge, value, size, edge_color in edges:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), value)
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), edge_color)
        borders.append(element)
    tbl_pr.append(borders)


def clear_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "none")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")
        borders.append(element)
    tbl_pr.append(borders)


def set_table_cell_margins(
    table, top: int = 70, bottom: int = 70, left: int = 110, right: int = 110
) -> None:
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, value in (
        ("top", top),
        ("left", left),
        ("bottom", bottom),
        ("right", right),
    ):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)


def set_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def set_table_width_percent(table, percent: int = 100) -> None:
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), str(percent * 50))


def set_column_widths(table, ratios: list[float]) -> None:
    total = sum(ratios)
    widths = [Cm(TEXT_WIDTH_CM * ratio / total) for ratio in ratios]
    for index, width in enumerate(widths):
        if index < len(table.columns):
            table.columns[index].width = width
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index < len(widths):
                cell.width = widths[index]


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_row_height(row, cm: float) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(int(cm * 567)))
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)


def add_page_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)
    style_run(run, size=8, color=GREY)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

# Word and WPS expect property elements in schema order. python-docx preserves
# order for its own API calls, but the manual XML helpers above append, so the
# whole package is normalised once before saving.
SCHEMA_ORDERS = {
    "pPr": [
        "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
        "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
        "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
        "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
        "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
        "suppressOverlap", "jc", "textDirection", "textAlignment",
        "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
        "pPrChange",
    ],
    "tblPr": [
        "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
        "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
        "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
        "tblCaption", "tblDescription",
    ],
    "trPr": [
        "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
        "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc",
        "hidden", "ins", "del", "trPrChange",
    ],
    "tcPr": [
        "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
        "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    ],
    "rPr": [
        "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
        "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
        "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
        "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
        "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
        "eastAsianLayout", "specVanish", "oMath",
    ],
}


def normalize_xml(root) -> None:
    for tag, order in SCHEMA_ORDERS.items():
        for element in root.iter(qn(f"w:{tag}")):
            children = list(element)
            children.sort(
                key=lambda child: order.index(child.tag.split("}")[-1])
                if child.tag.split("}")[-1] in order
                else len(order)
            )
            for child in list(element):
                element.remove(child)
            for child in children:
                element.append(child)

def add_inline(
    paragraph,
    text: str,
    size: float = 10.5,
    color: RGBColor = INK,
    bold: bool = False,
) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            style_run(
                paragraph.add_run(text[position : match.start()]),
                size=size,
                color=color,
                bold=bold,
            )
        token = match.group(0)
        if token.startswith("**"):
            style_run(
                paragraph.add_run(token[2:-2]), size=size, color=color, bold=True
            )
        else:
            style_run(
                paragraph.add_run(token[1:-1]),
                name="Consolas",
                size=size - 0.5,
                color=NAVY,
            )
        position = match.end()
    if position < len(text):
        style_run(
            paragraph.add_run(text[position:]), size=size, color=color, bold=bold
        )


# ---------------------------------------------------------------------------
# Markdown parsing (deliberately small: only the subset used in the source)
# ---------------------------------------------------------------------------

def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown(source: str) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    paragraph_lines: list[str] = []
    index = 0
    lines = source.splitlines()

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append(("p", " ".join(paragraph_lines)))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            blocks.append(("code", "\n".join(code_lines)))
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = []
            for row_index, row_line in enumerate(table_lines):
                cells = _split_table_row(row_line)
                if row_index == 1 and all(
                    re.fullmatch(r":?-{2,}:?", cell) for cell in cells
                ):
                    continue
                rows.append(cells)
            blocks.append(("table", rows))
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            blocks.append(("note", " ".join(quote_lines)))
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            blocks.append((f"h{level}", heading.group(2).strip()))
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            flush_paragraph()
            blocks.append(("bullet", bullet.group(1).strip()))
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            flush_paragraph()
            blocks.append(("num", (numbered.group(1), numbered.group(2).strip())))
            index += 1
            continue

        paragraph_lines.append(stripped)
        index += 1

    flush_paragraph()
    return blocks


# ---------------------------------------------------------------------------
# Document construction
# ---------------------------------------------------------------------------

def setup_page(section) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)


def setup_styles(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        rfonts.set(qn("w:eastAsia"), "Calibri")

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(10.5)
    bullet.font.color.rgb = INK
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.15

    heading_specs = {
        "Heading 1": (15, NAVY, 22, 8, True),
        "Heading 2": (11.5, MID_BLUE, 13, 4, False),
        "Heading 3": (10.5, INK, 10, 3, False),
    }
    for name, (size, color, before, after, border) in heading_specs.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style_rpr = style.element.get_or_add_rPr()
        style_rfonts = style_rpr.find(qn("w:rFonts"))
        if style_rfonts is not None:
            style_rfonts.set(qn("w:eastAsia"), "Calibri")
        if border:
            p_pr = style.element.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:space"), "6")
            bottom.set(qn("w:color"), NAVY_HEX)
            p_bdr.append(bottom)
            p_pr.append(p_bdr)


def setup_header_footer(section) -> None:
    section.different_first_page_header_footer = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.paragraph_format.tab_stops.add_tab_stop(
        USABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT
    )
    header_paragraph.paragraph_format.space_after = Pt(2)
    style_run(
        header_paragraph.add_run("AI CUSTOMER CALLING AGENT"),
        size=8,
        bold=True,
        color=NAVY,
    )
    style_run(
        header_paragraph.add_run("\tWeek 1  |  Discovery and Foundation"),
        size=8,
        color=GREY,
    )
    add_paragraph_border(header_paragraph, "bottom", color="C9D3E3", size=6, space=4)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(
        USABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT
    )
    footer_paragraph.paragraph_format.space_before = Pt(2)
    add_paragraph_border(footer_paragraph, "top", color="C9D3E3", size=6, space=4)
    style_run(
        footer_paragraph.add_run("Version 1.0  |  28 October 2026"),
        size=8,
        color=GREY,
    )
    style_run(footer_paragraph.add_run("\tPage "), size=8, color=GREY)
    add_page_field(footer_paragraph, "PAGE")
    style_run(footer_paragraph.add_run(" of "), size=8, color=GREY)
    add_page_field(footer_paragraph, "NUMPAGES")


def build_cover(document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width_percent(table)
    set_fixed_layout(table)
    clear_table_borders(table)
    set_table_cell_margins(table, top=200, bottom=200, left=420, right=420)
    set_column_widths(table, [100])

    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    shade_cell(cell, BLACK_HEX)
    set_row_height(table.rows[0], 10.6)

    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(26)
    style_run(first.add_run("INTERNSHIP PROJECT  /  WEEK 1"), size=9, bold=True, color=SOFT_BLUE)

    title_one = cell.add_paragraph()
    title_one.paragraph_format.space_after = Pt(0)
    title_one.paragraph_format.line_spacing = 1.0
    style_run(title_one.add_run("AI Customer"), size=31, bold=True, color=WHITE)

    title_two = cell.add_paragraph()
    title_two.paragraph_format.space_after = Pt(4)
    title_two.paragraph_format.line_spacing = 1.0
    style_run(title_two.add_run("Calling Agent"), size=31, bold=True, color=WHITE)

    divider = cell.add_paragraph()
    divider.paragraph_format.space_before = Pt(16)
    divider.paragraph_format.space_after = Pt(16)
    add_paragraph_border(divider, "bottom", color="2E5B96", size=12, space=1)

    subtitle = cell.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    style_run(
        subtitle.add_run("Discovery and Foundation"),
        size=16,
        bold=True,
        color=PALE_BLUE,
    )

    document_type = cell.add_paragraph()
    document_type.paragraph_format.space_after = Pt(24)
    style_run(
        document_type.add_run("Development Progress Document"),
        size=11,
        color=SOFT_BLUE,
    )

    status = cell.add_paragraph()
    status.paragraph_format.space_after = Pt(0)
    style_run(
        status.add_run(
            "Prototype implemented  |  27 unit tests passing  |  Live call pending access"
        ),
        size=9,
        italic=True,
        color=RGBColor(0x9F, 0xB3, 0xC8),
    )

    gap = document.add_paragraph()
    gap.paragraph_format.space_after = Pt(14)

    band = document.add_table(rows=1, cols=1)
    band.autofit = False
    set_table_width_percent(band)
    set_fixed_layout(band)
    clear_table_borders(band)
    set_table_cell_margins(band, top=90, bottom=90, left=140, right=140)
    set_column_widths(band, [100])
    band_cell = band.cell(0, 0)
    shade_cell(band_cell, NAVY_HEX)
    band_paragraph = band_cell.paragraphs[0]
    band_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    band_paragraph.paragraph_format.space_after = Pt(0)
    style_run(
        band_paragraph.add_run(
            "WEEK 1 OF 4   |   DAYS 1 TO 5   |   VERSION 1.0   |   28 OCTOBER 2026"
        ),
        size=8.5,
        bold=True,
        color=WHITE,
    )

    gap_two = document.add_paragraph()
    gap_two.paragraph_format.space_after = Pt(10)

    meta_rows = [
        ("PREPARED BY", META["prepared_by"]),
        ("PROJECT", META["project"]),
        ("DOCUMENT TYPE", META["document_type"]),
        ("WEEK COVERED", META["week"]),
        ("ENVIRONMENT", META["environment"]),
        ("RELATED PLAN", META["related"]),
        ("REPOSITORY STATE", META["repository"]),
    ]
    meta_table = document.add_table(rows=len(meta_rows), cols=2)
    meta_table.autofit = False
    set_table_width_percent(meta_table)
    set_fixed_layout(meta_table)
    clear_table_borders(meta_table)
    set_table_cell_margins(meta_table, top=55, bottom=55, left=0, right=80)
    set_column_widths(meta_table, [24, 76])
    for row_index, (label, value) in enumerate(meta_rows):
        label_cell, value_cell = meta_table.rows[row_index].cells
        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.paragraph_format.space_after = Pt(0)
        style_run(label_paragraph.add_run(label), size=7.5, bold=True, color=NAVY)
        value_paragraph = value_cell.paragraphs[0]
        value_paragraph.paragraph_format.space_after = Pt(0)
        style_run(value_paragraph.add_run(value), size=9.5, color=INK)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    style_run(
        note.add_run(
            "Prepared for internship supervisor review. This document reports implemented "
            "and verified work only. No live phone call has been placed."
        ),
        size=8,
        italic=True,
        color=GREY,
    )

    document.add_page_break()


def build_contents(document, entries: list[tuple[str, str]]) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(16)
    add_paragraph_border(heading, "bottom", color=NAVY_HEX, size=8, space=6)
    style_run(heading.add_run("Contents"), size=22, bold=True, color=INK)

    table = document.add_table(rows=len(entries), cols=2)
    table.autofit = False
    set_table_width_percent(table)
    set_fixed_layout(table)
    clear_table_borders(table)
    set_table_cell_margins(table, top=75, bottom=75, left=0, right=100)
    set_column_widths(table, [16, 84])

    for row_index, (number, title) in enumerate(entries):
        number_cell, title_cell = table.rows[row_index].cells
        number_paragraph = number_cell.paragraphs[0]
        number_paragraph.paragraph_format.space_after = Pt(0)
        number_paragraph.paragraph_format.line_spacing = 1.0
        style_run(number_paragraph.add_run(number), size=9.5, bold=True, color=NAVY)

        title_paragraph = title_cell.paragraphs[0]
        title_paragraph.paragraph_format.space_after = Pt(0)
        title_paragraph.paragraph_format.line_spacing = 1.0
        style_run(title_paragraph.add_run(title), size=10, color=INK)
        add_paragraph_border(title_paragraph, "bottom", color="E4E9F1", size=4, space=3)

    document.add_page_break()


def render_heading(document, text: str, level: int) -> None:
    style_name = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
    document.add_paragraph(text, style=style_name)


def render_paragraph(document, text: str) -> None:
    paragraph = document.add_paragraph()
    add_inline(paragraph, text)


def render_bullet(document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    add_inline(paragraph, text)


def render_numbered(document, number: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.first_line_indent = Cm(-0.55)
    paragraph.paragraph_format.space_after = Pt(3)
    style_run(paragraph.add_run(f"{number}."), size=10.5, bold=True, color=NAVY)
    style_run(paragraph.add_run("  "), size=10.5)
    add_inline(paragraph, text)


def render_note(document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_width_percent(table)
    set_fixed_layout(table)
    clear_table_borders(table)
    set_table_cell_margins(table, top=130, bottom=130, left=220, right=180)
    set_column_widths(table, [100])
    cell = table.cell(0, 0)
    shade_cell(cell, NOTE_HEX)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, text, size=9.5, color=INK)
    add_paragraph_border(paragraph, "left", color=NAVY_HEX, size=24, space=8)

    gap = document.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    style_run(gap.add_run(""), size=4)


def render_code(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.left_indent = Cm(0.15)
    shade_paragraph(paragraph, CODE_HEX)
    add_paragraph_border(paragraph, "left", color=NAVY_HEX, size=18, space=6)

    for line_index, line in enumerate(text.split("\n")):
        if line_index:
            paragraph.add_run().add_break()
        style_run(
            paragraph.add_run(line),
            name="Consolas",
            size=8.5,
            color=RGBColor(0x1B, 0x1F, 0x24),
        )


def render_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    header = rows[0]
    column_count = len(header)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width_percent(table)
    set_fixed_layout(table)
    set_table_borders(table)
    set_table_cell_margins(table)
    ratios = WIDTH_MAP.get(tuple(header), DEFAULT_RATIOS.get(column_count, [1] * column_count))
    set_column_widths(table, ratios)

    for row_index, row_cells in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            text = row_cells[column_index] if column_index < len(row_cells) else ""
            if row_index == 0:
                shade_cell(cell, NAVY_HEX)
                add_inline(paragraph, text, size=8.6, color=WHITE, bold=True)
            else:
                if row_index % 2 == 0:
                    shade_cell(cell, CARD_HEX)
                add_inline(paragraph, text, size=8.8, color=INK)

        if row_index == 0:
            repeat_header(table.rows[row_index])
        prevent_row_split(table.rows[row_index])

    gap = document.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    style_run(gap.add_run(""), size=4)


def render_blocks(document, blocks: list[tuple[str, object]]) -> None:
    for kind, payload in blocks:
        if kind == "h1":
            continue
        if kind == "h2":
            render_heading(document, payload, 2)
        elif kind == "h3":
            render_heading(document, payload, 3)
        elif kind == "h4":
            render_heading(document, payload, 4)
        elif kind == "p":
            render_paragraph(document, payload)
        elif kind == "bullet":
            render_bullet(document, payload)
        elif kind == "num":
            number, text = payload
            render_numbered(document, number, text)
        elif kind == "note":
            render_note(document, payload)
        elif kind == "code":
            render_code(document, payload)
        elif kind == "table":
            render_table(document, payload)


def collect_contents(blocks: list[tuple[str, object]]) -> list[tuple[str, str]]:
    entries: list[tuple[str, str]] = []
    for kind, payload in blocks:
        if kind != "h2":
            continue
        text = str(payload)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", text)
        appendix = re.match(r"^(Appendix\s+[A-D])\s*[-–]\s*(.*)$", text)
        if numbered:
            entries.append((f"{numbered.group(1)}.", numbered.group(2)))
        elif appendix:
            entries.append((appendix.group(1), appendix.group(2)))
    return entries


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Markdown source not found: {SOURCE}")

    blocks = parse_markdown(SOURCE.read_text(encoding="utf-8"))
    contents = collect_contents(blocks)

    document = Document()
    setup_page(document.sections[0])
    setup_styles(document)
    setup_header_footer(document.sections[0])

    properties = document.core_properties
    properties.title = "AI Customer Calling Agent - Week 1 Development Document"
    properties.author = "Sumyai Butt"
    properties.subject = "Week 1: Discovery and Foundation"
    properties.keywords = "AI calling agent, customer support, Week 1, prototype, Python"
    properties.category = "Internship project documentation"
    properties.comments = (
        "Week 1 prototype implemented and unit tested. No live call has been placed."
    )

    build_cover(document)
    build_contents(document, contents)
    render_blocks(document, blocks)

    section = document.sections[0]
    for root in (
        document.element,
        document.styles.element,
        section.header._element,
        section.footer._element,
    ):
        normalize_xml(root)

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)

    document.save(OUTPUT)
    print(f"Built {OUTPUT}")
    print(f"Sections: {len(contents)} | Tables: {len(document.tables)} | Paragraphs: {len(document.paragraphs)}")


if __name__ == "__main__":
    main()
